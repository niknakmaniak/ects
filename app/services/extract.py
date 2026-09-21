import json
import logging
from pathlib import Path

from pypdf import PdfReader
from docx import Document
from pptx import Presentation

logger = logging.getLogger(__name__)


def _read_text_file(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        parts.append(f"--- page {i} ---\n{text}")
    return "\n\n".join(parts)


def extract_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_pptx(path: Path) -> str:
    prs = Presentation(str(path))
    chunks = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        chunks.append(f"--- slide {i} ---\n" + "\n".join(texts))
    return "\n\n".join(chunks)


def classify_file(path: Path) -> str:
    name = path.name.lower()
    if path.suffix.lower() in {".m4a", ".mp3", ".wav"}:
        return "audio"
    if "transcript" in name or path.suffix.lower() in {".vtt", ".srt"}:
        return "audio"
    if "syllabus" in name or "notes" in name or path.suffix.lower() == ".docx" and "support" not in name:
        return "syllabus"
    if path.suffix.lower() in {".ppt", ".pptx"} or "support" in name:
        return "support"
    if path.suffix.lower() == ".pdf":
        return "support"
    if path.suffix.lower() == ".txt":
        return "audio"
    return "other"


def extract_all(job, work: Path) -> dict:
    input_dir = work / "input"
    out_dir = work / "extracted"
    out_dir.mkdir(parents=True, exist_ok=True)
    bundle = {"syllabus": "", "support": "", "audio": "", "files": []}

    for path in sorted(input_dir.rglob("*")):
        if not path.is_file():
            continue
        if path.name.lower() == "ready.txt":
            continue
        kind = classify_file(path)
        text = ""
        try:
            if path.suffix.lower() == ".pdf":
                text = extract_pdf(path)
            elif path.suffix.lower() in {".docx", ".doc"}:
                text = extract_docx(path)
            elif path.suffix.lower() in {".pptx", ".ppt"}:
                text = extract_pptx(path)
            elif path.suffix.lower() in {".txt", ".vtt", ".srt", ".md"}:
                text = _read_text_file(path)
        except Exception as e:
            logger.warning("Extraction ignoree pour %s: %s", path.name, e)
            bundle["files"].append({"path": str(path.relative_to(work)), "kind": kind, "error": str(e)})
            continue

        rel = str(path.relative_to(work))
        bundle["files"].append({"path": rel, "kind": kind, "chars": len(text)})
        if kind == "syllabus" and text:
            bundle["syllabus"] += f"\n\n=== {path.name} ===\n{text}"
        elif kind == "support" and text:
            bundle["support"] += f"\n\n=== {path.name} ===\n{text}"
        elif kind == "audio" and text:
            bundle["audio"] += f"\n\n=== {path.name} ===\n{text}"

        if text:
            (out_dir / f"{path.stem}.txt").write_text(text, encoding="utf-8")

    (out_dir / "bundle.json").write_text(json.dumps(bundle, indent=2, ensure_ascii=False), encoding="utf-8")
    return bundle
